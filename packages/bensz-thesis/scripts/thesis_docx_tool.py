#!/usr/bin/env python3
"""Generic DOCX exporter for bensz-thesis projects.

The exporter intentionally produces an editable Word draft, not a
pixel-perfect clone of the PDF. It keeps PDF build behavior untouched and
uses a source-based conversion path:

    thesis LaTeX -> normalized Markdown -> HTML5 + MathML -> DOCX

Complex LaTeX objects are degraded to explicit placeholders and the original
source is stored under ``.latex-cache/docx/unsupported/`` for manual cleanup.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable

try:
    from docx import Document as DocxDocument  # type: ignore
except Exception:
    DocxDocument = None


CACHE_DIRNAME = ".latex-cache"
DOCX_CACHE_DIRNAME = "docx"
SKIP_SOURCE_NAMES = {
    "@config.tex",
    "config-pre.tex",
    "config.tex",
    "preamble.tex",
    "packages.tex",
}
GRAPHICS_SUFFIXES = (
    "",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".webp",
    ".svg",
    ".pdf",
)
PANDOC_CANDIDATES = (
    "/opt/homebrew/bin/pandoc",
    "/usr/local/bin/pandoc",
)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = {"w": W_NS}


class DocxExportError(RuntimeError):
    """Raised when DOCX export cannot continue."""


class TexSource:
    """A collected TeX source file and its raw text."""

    def __init__(self, path: Path, text: str) -> None:
        self.path = path
        self.text = text


class DocxExportState:
    """Mutable export facts used by conversion and quality reporting."""

    def __init__(self, cache_dir: Path) -> None:
        self.cache_dir = cache_dir
        self.unsupported_dir = cache_dir / "unsupported"
        self.included_sources: list[Path] = []
        self.missing_sources: list[str] = []
        self.skipped_files: list[str] = []
        self.missing_assets: list[str] = []
        self.bibliography_files: list[Path] = []
        self.unsupported_counts: Counter[str] = Counter()
        self.heading_counts: Counter[int] = Counter()
        self.warnings: list[str] = []
        self.fallback_used = False

    def add_source(self, path: Path) -> None:
        if path not in self.included_sources:
            self.included_sources.append(path)

    def add_bibliography(self, path: Path) -> None:
        resolved = path.resolve()
        if resolved not in self.bibliography_files:
            self.bibliography_files.append(resolved)

    def add_missing_asset(self, asset: str) -> None:
        if asset not in self.missing_assets:
            self.missing_assets.append(asset)

    def add_warning(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)

    def add_unsupported(self, env_name: str, original_source: str, caption: str = "") -> str:
        self.unsupported_counts[env_name] += 1
        index = self.unsupported_counts[env_name]
        self.unsupported_dir.mkdir(parents=True, exist_ok=True)
        saved_path = self.unsupported_dir / f"{env_name}-{index:03d}.tex"
        saved_path.write_text(original_source.strip() + "\n", encoding="utf-8")
        label = caption or "内容略"
        return (
            f"\n\n> {env_name}：{label}（需在 Word 中人工整理；"
            f"LaTeX 源已保存到 {saved_path.relative_to(self.cache_dir)}）\n\n"
        )


def resolve_executable(name: str) -> str:
    resolved = shutil.which(name)
    if resolved:
        return resolved
    if name == "pandoc":
        for candidate in PANDOC_CANDIDATES:
            if Path(candidate).exists():
                return candidate
    raise FileNotFoundError(f"未找到可执行文件：{name}")


def run_cmd(
    args: list[str],
    *,
    cwd: Path | None = None,
    input_text: str | None = None,
) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        args,
        cwd=cwd,
        input=input_text,
        text=True,
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        raise DocxExportError(
            "命令执行失败：{}\nstdout:\n{}\nstderr:\n{}".format(
                " ".join(args),
                result.stdout,
                result.stderr,
            )
        )
    return result


def get_pandoc_version() -> str:
    try:
        result = run_cmd([resolve_executable("pandoc"), "--version"])
    except Exception as exc:
        return f"unavailable: {exc}"
    first = result.stdout.splitlines()[0] if result.stdout.splitlines() else "pandoc"
    return first.strip()


def strip_comments(text: str) -> str:
    lines: list[str] = []
    for line in text.splitlines():
        lines.append(re.sub(r"(?<!\\)%.*$", "", line))
    return "\n".join(lines)


def parse_braced(text: str, start: int) -> tuple[str, int]:
    if start >= len(text) or text[start] != "{":
        raise ValueError("parse_braced expects '{' at start")
    depth = 0
    out: list[str] = []
    i = start
    while i < len(text):
        ch = text[i]
        if ch == "\\":
            if i + 1 < len(text):
                out.append(text[i : i + 2])
                i += 2
                continue
            out.append(ch)
            i += 1
            continue
        if ch == "{":
            depth += 1
            if depth > 1:
                out.append(ch)
            i += 1
            continue
        if ch == "}":
            depth -= 1
            if depth == 0:
                return "".join(out), i + 1
            out.append(ch)
            i += 1
            continue
        out.append(ch)
        i += 1
    raise ValueError("unbalanced braces")


def flatten_texorpdfstring(text: str) -> str:
    pattern = re.compile(r"\\texorpdfstring\s*\{")
    while True:
        match = pattern.search(text)
        if not match:
            return text
        try:
            first_end = parse_braced(text, match.end() - 1)[1]
            i = first_end
            while i < len(text) and text[i].isspace():
                i += 1
            if i >= len(text) or text[i] != "{":
                return text
            second, second_end = parse_braced(text, i)
        except Exception:
            return text
        text = text[: match.start()] + second + text[second_end:]


def _extract_braced_command_args(text: str, start: int, n_args: int) -> tuple[list[str], int] | None:
    i = start
    args: list[str] = []
    for _ in range(n_args):
        while i < len(text) and text[i].isspace():
            i += 1
        if i >= len(text) or text[i] != "{":
            return None
        value, i = parse_braced(text, i)
        args.append(value)
    return args, i


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _with_tex_suffix(raw: str) -> str:
    path = Path(raw.strip())
    if path.suffix:
        return path.as_posix()
    return path.with_suffix(".tex").as_posix()


def _is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _resolve_include_path(
    raw: str,
    current_dir: Path,
    project_dir: Path,
    state: DocxExportState | None = None,
) -> Path | None:
    rel = Path(_with_tex_suffix(raw))
    candidates = []
    if rel.is_absolute():
        candidates.append(rel)
    else:
        candidates.extend([project_dir / rel, current_dir / rel])
    for candidate in candidates:
        resolved = candidate.resolve()
        if not resolved.exists():
            continue
        if _is_within(resolved, project_dir):
            return resolved
        if state is not None:
            state.add_warning(f"Blocked external TeX input: {resolved}")
        return None
    return None


def _display_missing_include(raw: str) -> str:
    return _with_tex_suffix(raw)


def _iter_input_matches(text: str) -> Iterable[re.Match[str]]:
    cleaned = strip_comments(text)
    return re.finditer(r"\\(?:input|include)\s*\{([^{}]+)\}", cleaned)


def _is_match_in_comment(text: str, start: int) -> bool:
    line_start = text.rfind("\n", 0, start) + 1
    prefix = text[line_start:start]
    return re.search(r"(?<!\\)%", prefix) is not None


def _iter_input_matches_with_original_offsets(text: str) -> Iterable[re.Match[str]]:
    for match in re.finditer(r"\\(?:input|include)\s*\{([^{}]+)\}", text):
        if _is_match_in_comment(text, match.start()):
            continue
        yield match


def collect_tex_sources(
    tex_file: Path,
    project_dir: Path,
    state: DocxExportState,
    _visited: set[Path] | None = None,
) -> list[TexSource]:
    visited = _visited if _visited is not None else set()
    tex_path = tex_file.resolve()
    if tex_path in visited:
        return []
    visited.add(tex_path)
    text = _read_text(tex_path)
    state.add_source(tex_path)
    collected = [TexSource(tex_path, text)]
    for match in _iter_input_matches(text):
        raw = match.group(1).strip()
        if Path(raw).name in SKIP_SOURCE_NAMES:
            if raw not in state.skipped_files:
                state.skipped_files.append(raw)
            continue
        child = _resolve_include_path(raw, tex_path.parent, project_dir, state)
        if child is None:
            missing = _display_missing_include(raw)
            if missing not in state.missing_sources:
                state.missing_sources.append(missing)
            continue
        collected.extend(collect_tex_sources(child, project_dir, state, visited))
    return collected


def parse_graphicspaths(text: str) -> list[Path]:
    paths: list[Path] = []
    cleaned = strip_comments(text)
    for match in re.finditer(r"\\graphicspath\s*\{", cleaned):
        try:
            payload, _ = parse_braced(cleaned, match.end() - 1)
        except Exception:
            continue
        for inner in re.finditer(r"\{([^{}]+)\}", payload):
            value = inner.group(1).strip()
            if value:
                paths.append(Path(value))
    return paths


def build_graphic_search_dirs(project_dir: Path, main_text: str) -> list[Path]:
    dirs = [Path(".")]
    dirs.extend(parse_graphicspaths(main_text))
    for name in ("assets", "figures", "images"):
        if (project_dir / name).exists():
            dirs.append(Path(name))
    deduped: list[Path] = []
    seen: set[str] = set()
    for item in dirs:
        key = item.as_posix()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def resolve_graphics_path(
    image_ref: str,
    project_dir: Path,
    source_dir: Path,
    graphic_search_dirs: Iterable[Path],
    state: DocxExportState,
) -> str:
    raw = image_ref.strip()
    if not raw or "://" in raw:
        return raw
    image_path = Path(raw)
    if image_path.is_absolute():
        if not image_path.exists():
            state.add_missing_asset(raw)
        elif not _is_within(image_path, project_dir):
            state.add_missing_asset(raw)
            state.add_warning(f"Blocked external graphics asset: {image_path.resolve()}")
            return ""
        return raw
    try:
        source_rel = source_dir.resolve().relative_to(project_dir.resolve())
    except ValueError:
        source_rel = Path(".")
    roots = [Path("."), source_rel, *graphic_search_dirs]
    seen: set[str] = set()
    for root in roots:
        key = root.as_posix()
        if key in seen:
            continue
        seen.add(key)
        for suffix in GRAPHICS_SUFFIXES:
            candidate = root / image_path
            if suffix:
                candidate = candidate.with_suffix(suffix)
            absolute_candidate = (project_dir / candidate).resolve()
            if not absolute_candidate.exists():
                continue
            if _is_within(absolute_candidate, project_dir):
                return candidate.as_posix()
            state.add_missing_asset(raw)
            state.add_warning(f"Blocked external graphics asset: {absolute_candidate}")
            return ""
    state.add_missing_asset(raw)
    return raw


def detect_bibliography_files(text: str, project_dir: Path, state: DocxExportState) -> None:
    cleaned = strip_comments(text)
    for match in re.finditer(r"\\addbibresource(?:\[[^\]]*\])?\{([^{}]+)\}", cleaned):
        bib = project_dir / match.group(1).strip()
        if bib.exists():
            state.add_bibliography(bib)
        else:
            state.add_warning(f"Missing bibliography: {match.group(1).strip()}")
    for match in re.finditer(r"\\bibliography\{([^{}]+)\}", cleaned):
        for raw in match.group(1).split(","):
            name = raw.strip()
            if not name:
                continue
            path = Path(name)
            if not path.suffix:
                path = path.with_suffix(".bib")
            bib = project_dir / path
            if bib.exists():
                state.add_bibliography(bib)
            else:
                state.add_warning(f"Missing bibliography: {path.as_posix()}")


def _protect_inline_math(text: str) -> tuple[str, dict[str, str]]:
    protected: dict[str, str] = {}

    def repl(match: re.Match[str]) -> str:
        token = f"@@MATH{len(protected)}@@"
        protected[token] = match.group(0)
        return token

    return re.sub(r"(?<!\\)\$(?!\$).*?(?<!\\)\$", repl, text), protected


def _restore_tokens(text: str, tokens: dict[str, str]) -> str:
    for token, value in tokens.items():
        text = text.replace(token, value)
    return text


def _replace_simple_text_commands(text: str) -> str:
    replacements = {
        "textbf": r"**\1**",
        "bfseries": r"\1",
        "textit": r"*\1*",
        "emph": r"*\1*",
        "underline": r"\1",
        "texttt": r"`\1`",
        "mbox": r"\1",
        "textrm": r"\1",
        "textnormal": r"\1",
    }
    changed = True
    while changed:
        before = text
        for cmd, repl in replacements.items():
            text = re.sub(r"\\" + cmd + r"\{([^{}]*)\}", repl, text)
        changed = before != text
    return text


def latex_inline_to_text(text: str) -> str:
    text, math_tokens = _protect_inline_math(text)
    text = _replace_simple_text_commands(text)
    text = text.replace("\\LaTeX", "LaTeX")
    text = text.replace("\\quad", " ")
    text = text.replace("\\qquad", " ")
    text = text.replace("\\~{}", "~")
    text = text.replace("\\%", "%")
    text = re.sub(r"\\url\{([^{}]*)\}", r"<\1>", text)
    text = re.sub(r"\\path\{([^{}]*)\}", r"`\1`", text)
    text = re.sub(r"\\verb(.)(.*?)\1", r"`\2`", text)
    text = re.sub(r"\\label\{[^{}]*\}", "", text)

    def cite_repl(match: re.Match[str]) -> str:
        keys = [item.strip() for item in match.group(2).split(",") if item.strip()]
        return "[" + "; ".join(f"@{key}" for key in keys) + "]" if keys else ""

    text = re.sub(
        r"\\(?:cite|citep|citet|parencite|textcite|autocite|supercite)(?:\[[^\]]*\]){0,2}\{([^{}]+)\}",
        lambda m: "[" + "; ".join(f"@{k.strip()}" for k in m.group(1).split(",") if k.strip()) + "]",
        text,
    )
    text = re.sub(
        r"\\(?:ref|eqref|autoref|cref|Cref)\{([^{}]+)\}",
        lambda m: f"`{m.group(1).strip()}`",
        text,
    )
    text = re.sub(r"\\(?:gls|Gls|acrshort|acrlong|acrlongpl)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\[a-zA-Z@]+\*?(?:\[[^\]]*\])?", "", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = _restore_tokens(text, math_tokens)
    return text.strip()


def _caption_from_block(block: str) -> str:
    for cmd in ("bicaption", "caption"):
        pattern = re.compile(r"\\" + cmd + r"(?:\[[^\]]*\])?\s*\{")
        match = pattern.search(block)
        if not match:
            continue
        try:
            first, end = parse_braced(block, match.end() - 1)
            if cmd == "bicaption":
                i = end
                while i < len(block) and block[i].isspace():
                    i += 1
                if i < len(block) and block[i] == "{":
                    second, _ = parse_braced(block, i)
                    return latex_inline_to_text(first or second)
            return latex_inline_to_text(first)
        except Exception:
            continue
    return ""


def _replace_env(text: str, env_name: str, replacer) -> str:
    pattern = re.compile(
        r"\\begin\{" + re.escape(env_name) + r"\}(?:\[[^\]]*\])?(.*?)\\end\{" + re.escape(env_name) + r"\}",
        re.S,
    )
    return pattern.sub(lambda match: replacer(match.group(1), match.group(0)), text)


def _replace_math_envs(text: str) -> str:
    def repl(block: str, _original: str) -> str:
        cleaned = re.sub(r"\\label\{[^{}]*\}", "", block).strip()
        return f"\n\n$$\n{cleaned}\n$$\n\n" if cleaned else ""

    for env in ("equation", "equation*", "align", "align*", "gather", "gather*", "multline", "multline*"):
        text = _replace_env(text, env, repl)
    return text


def _replace_headings(text: str, state: DocxExportState) -> str:
    levels = {
        "chapter": 1,
        "section": 2,
        "subsection": 3,
        "subsubsection": 4,
        "paragraph": 5,
    }
    pattern = re.compile(r"\\(chapter|section|subsection|subsubsection|paragraph)(\*)?(?:\[[^\]]*\])?\s*\{")
    out: list[str] = []
    cursor = 0
    while True:
        match = pattern.search(text, cursor)
        if not match:
            out.append(text[cursor:])
            break
        out.append(text[cursor : match.start()])
        try:
            title, end = parse_braced(text, match.end() - 1)
        except Exception:
            out.append(text[match.start() : match.end()])
            cursor = match.end()
            continue
        level = levels[match.group(1)]
        if match.group(2) and match.group(1) == "chapter":
            level = 1
        state.heading_counts[level] += 1
        hashes = "#" * min(level, 6)
        out.append(f"\n\n{hashes} {latex_inline_to_text(title)}\n\n")
        cursor = end
    return "".join(out)


def convert_latex_to_markdown(
    text: str,
    *,
    project_dir: Path,
    source_dir: Path,
    state: DocxExportState,
    graphic_search_dirs: Iterable[Path],
) -> str:
    text = strip_comments(text)
    text = flatten_texorpdfstring(text)
    detect_bibliography_files(text, project_dir, state)

    text = re.sub(r"\\begin\{document\}|\\end\{document\}", "", text)
    text = re.sub(r"\\(?:clearpage|cleardoublepage|newpage|frontmatter|mainmatter|backmatter)\b", "\n\n", text)
    text = re.sub(r"\\(?:pagestyle|thispagestyle|pagenumbering|setcounter|addcontentsline|geometry|newgeometry)\s*(?:\[[^\]]*\])?(?:\{[^{}]*\}){0,3}", "", text)
    text = text.replace("\\appendix", "\n\n# 附录\n\n")
    text = re.sub(r"\\tableofcontents\b", "\n\n# 目录\n\n> 目录域请在 Word 中右键更新。\n\n", text)
    text = re.sub(r"\\listoffigures\b", "\n\n# 图目录\n\n> 图目录请在 Word 中右键更新。\n\n", text)
    text = re.sub(r"\\listoftables\b", "\n\n# 表目录\n\n> 表目录请在 Word 中右键更新。\n\n", text)
    text = re.sub(r"\\listofmaterials(?:\[[^\]]*\])?", "\n\n# 图表目录\n\n> 图表目录请在 Word 中右键更新。\n\n", text)
    text = re.sub(r"\\listofnotations\b", "\n\n# 符号说明\n\n> 符号说明请在 Word 中检查更新。\n\n", text)
    text = re.sub(r"\\printbibliography(?:\[[^\]]*\])?", "\n\n# 参考文献\n\n> 参考文献请在 Word 中检查并按模板刷新。\n\n", text)
    text = re.sub(r"\\bibliography\{[^{}]+\}", "\n\n# 参考文献\n\n> 参考文献请在 Word 中检查并按模板刷新。\n\n", text)
    text = re.sub(r"\\nocite\{[^{}]*\}", "", text)

    text = _replace_math_envs(text)

    def figure_repl(block: str, original: str) -> str:
        caption = _caption_from_block(block)
        match = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^{}]+)\}", block)
        if not match:
            return state.add_unsupported("figure", original, caption or "未检测到图片")
        image = resolve_graphics_path(match.group(1), project_dir, source_dir, graphic_search_dirs, state)
        return f"\n\n![{caption}]({image})\n\n"

    text = _replace_env(text, "figure", figure_repl)

    def unsupported_repl(env_name: str):
        return lambda block, original: state.add_unsupported(env_name, original, _caption_from_block(block))

    for env_name in (
        "table",
        "longtable",
        "tabular",
        "tabularx",
        "algorithm",
        "algorithmic",
        "listing",
        "lstlisting",
        "minted",
        "tikzpicture",
    ):
        text = _replace_env(text, env_name, unsupported_repl(env_name))

    def inline_graphics(match: re.Match[str]) -> str:
        image = resolve_graphics_path(match.group(1), project_dir, source_dir, graphic_search_dirs, state)
        return f"\n\n![]({image})\n\n"

    text = re.sub(r"\\includegraphics(?:\[[^\]]*\])?\{([^{}]+)\}", inline_graphics, text)
    text = _replace_headings(text, state)
    text = re.sub(r"\\begin\{itemize\}|\\end\{itemize\}", "\n", text)
    text = re.sub(r"\\begin\{enumerate\}|\\end\{enumerate\}", "\n", text)
    text = re.sub(r"^[ \t]*\\item(?:\[[^\]]+\])?[ \t]*", "- ", text, flags=re.M)
    text = re.sub(r"\\inputminted(?:\[[^\]]*\])?\{([^{}]+)\}\{([^{}]+)\}", r"\n\n```\n[代码来自 \2]\n```\n\n", text)
    text = re.sub(r"\\mintinline\{[^{}]+\}\{([^{}]*)\}", r"`\1`", text)
    text = text.replace("\\\\", "\n")

    lines: list[str] = []
    in_math = False
    in_code = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            lines.append(stripped)
            continue
        if stripped == "$$":
            in_math = not in_math
            lines.append(stripped)
            continue
        if in_math or in_code or stripped.startswith("#") or stripped.startswith(">") or stripped.startswith("!["):
            lines.append(line.rstrip())
            continue
        if stripped.startswith("- "):
            lines.append("- " + latex_inline_to_text(stripped[2:]))
            continue
        cleaned = latex_inline_to_text(line)
        if cleaned:
            lines.append(cleaned)
        else:
            lines.append("")
    markdown = "\n".join(lines)
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)
    return markdown.strip() + "\n" if markdown.strip() else ""


def _document_body(text: str) -> str:
    begin = re.search(r"\\begin\{document\}", text)
    end = re.search(r"\\end\{document\}", text)
    if begin and end and begin.end() < end.start():
        return text[begin.end() : end.start()]
    return text


def _render_text_with_inputs(
    text: str,
    current_file: Path,
    project_dir: Path,
    state: DocxExportState,
    graphic_search_dirs: Iterable[Path],
    visited: set[Path],
) -> str:
    chunks: list[str] = []
    cursor = 0
    for match in _iter_input_matches_with_original_offsets(text):
        segment = text[cursor : match.start()]
        converted = convert_latex_to_markdown(
            segment,
            project_dir=project_dir,
            source_dir=current_file.parent,
            state=state,
            graphic_search_dirs=graphic_search_dirs,
        )
        if converted:
            chunks.append(converted)
        raw = match.group(1).strip()
        if Path(raw).name in SKIP_SOURCE_NAMES:
            if raw not in state.skipped_files:
                state.skipped_files.append(raw)
            cursor = match.end()
            continue
        child = _resolve_include_path(raw, current_file.parent, project_dir, state)
        if child is None:
            missing = _display_missing_include(raw)
            if missing not in state.missing_sources:
                state.missing_sources.append(missing)
            chunks.append(f"\n\n> 警告：未找到 TeX 文件 `{missing}`。\n\n")
            cursor = match.end()
            continue
        chunks.append(_render_file(child, project_dir, state, graphic_search_dirs, visited))
        cursor = match.end()
    tail = text[cursor:]
    converted_tail = convert_latex_to_markdown(
        tail,
        project_dir=project_dir,
        source_dir=current_file.parent,
        state=state,
        graphic_search_dirs=graphic_search_dirs,
    )
    if converted_tail:
        chunks.append(converted_tail)
    return "\n".join(chunk.strip() for chunk in chunks if chunk.strip()) + "\n"


def _render_file(
    tex_file: Path,
    project_dir: Path,
    state: DocxExportState,
    graphic_search_dirs: Iterable[Path],
    visited: set[Path],
) -> str:
    tex_path = tex_file.resolve()
    if tex_path in visited:
        return ""
    visited.add(tex_path)
    state.add_source(tex_path)
    text = _read_text(tex_path)
    return _render_text_with_inputs(text, tex_path, project_dir, state, graphic_search_dirs, visited)


def render_markdown(project_dir: Path, tex_file: Path, state: DocxExportState) -> str:
    main_text = _read_text(tex_file)
    detect_bibliography_files(main_text, project_dir, state)
    graphic_search_dirs = build_graphic_search_dirs(project_dir, main_text)
    body = _document_body(main_text)
    chunks = [
        "# 学位论文（LaTeX 源转换稿）",
        "",
        "> 本文件由 LaTeX 源自动转换为 editable Word draft；复杂对象会以占位符标出并进入质量报告。",
        "",
        _render_text_with_inputs(body, tex_file, project_dir, state, graphic_search_dirs, set()),
    ]
    for missing in state.missing_sources:
        state.add_warning(f"Missing TeX source: {missing}")
    return "\n".join(str(chunk).strip() for chunk in chunks if str(chunk).strip()) + "\n"


def discover_reference_doc(project_dir: Path, explicit: Path | None) -> Path | None:
    if explicit is not None:
        raw = explicit.expanduser()
        candidates = [raw]
        if not raw.is_absolute():
            candidates.append(project_dir / raw)
        checked: list[Path] = []
        for candidate in candidates:
            resolved = candidate.resolve()
            checked.append(resolved)
            if resolved.exists():
                return resolved
        raise FileNotFoundError(
            "--reference-doc 指向的文件不存在。已检查: "
            + ", ".join(str(item) for item in checked)
        )
    candidates = [project_dir / "artifacts" / "reference.docx"]
    candidates.extend(sorted((project_dir / "docs" / "official").glob("*.docx")))
    candidates.extend(sorted((project_dir / "docs").glob("*.docx")))
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    return None


def _resource_path(project_dir: Path) -> str:
    parts = [str(project_dir)]
    for name in ("assets", "figures", "images"):
        candidate = project_dir / name
        if candidate.exists():
            parts.append(str(candidate))
    return os.pathsep.join(parts)


def _copy_sanitized_bibliographies(bibliography: Iterable[Path], tmp_root: Path) -> list[Path]:
    copied: list[Path] = []
    for bib in bibliography:
        if not bib.exists():
            continue
        target = tmp_root / bib.name
        target.write_text(strip_comments(_read_text(bib)), encoding="utf-8")
        copied.append(target)
    return copied


def build_docx_from_markdown(
    markdown_path: Path,
    docx_path: Path,
    *,
    project_dir: Path,
    reference_doc: Path | None,
    bibliography: Iterable[Path],
    state: DocxExportState,
) -> None:
    pandoc = resolve_executable("pandoc")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="thesis-docx-", dir=str(state.cache_dir)) as tmp_dir:
        tmp_root = Path(tmp_dir)
        html_path = tmp_root / "main.html"
        sanitized_bibs = _copy_sanitized_bibliographies(bibliography, tmp_root)
        html_cmd = [
            pandoc,
            str(markdown_path),
            "-f",
            "markdown+raw_html+tex_math_dollars+superscript",
            "-t",
            "html5",
            "--mathml",
            f"--resource-path={_resource_path(project_dir)}",
            "-o",
            str(html_path),
        ]
        if sanitized_bibs:
            html_cmd.append("--citeproc")
            for bib in sanitized_bibs:
                html_cmd.extend(["--bibliography", str(bib)])
        direct_cmd = [
            pandoc,
            str(markdown_path),
            "--from=markdown+tex_math_dollars+raw_tex",
            "--to=docx",
            "--standalone",
            f"--resource-path={_resource_path(project_dir)}",
            "-o",
            str(docx_path),
        ]
        if reference_doc is not None:
            direct_cmd.append(f"--reference-doc={reference_doc}")
        for bib in sanitized_bibs:
            direct_cmd.extend(["--bibliography", str(bib)])
        if sanitized_bibs:
            direct_cmd.append("--citeproc")
        try:
            run_cmd(html_cmd, cwd=project_dir)
            docx_cmd = [
                pandoc,
                str(html_path),
                "-f",
                "html",
                f"--resource-path={_resource_path(project_dir)}",
                "-o",
                str(docx_path),
            ]
            if reference_doc is not None:
                docx_cmd.append(f"--reference-doc={reference_doc}")
            run_cmd(docx_cmd, cwd=project_dir)
        except Exception as html_exc:
            state.fallback_used = True
            state.add_warning(f"HTML5+MathML route failed; used direct Markdown->DOCX fallback: {html_exc}")
            run_cmd(direct_cmd, cwd=project_dir)


def _w_attr(key: str) -> str:
    return f"{{{W_NS}}}{key}"


def analyze_docx_styles(docx_path: Path) -> dict[str, object]:
    style_name_by_id: dict[str, str] = {}
    style_name_to_id: dict[str, str] = {}
    style_ids: set[str] = set()
    usage: Counter[str] = Counter()
    unknown: Counter[str] = Counter()
    if not docx_path.exists():
        return {
            "style_ids": style_ids,
            "style_name_by_id": style_name_by_id,
            "style_name_to_id": style_name_to_id,
            "usage": usage,
            "unknown": unknown,
        }
    with zipfile.ZipFile(docx_path, "r") as zf:
        if "word/styles.xml" in zf.namelist():
            root = ET.fromstring(zf.read("word/styles.xml"))
            for style in root.findall("w:style", W):
                if style.get(_w_attr("type"), "") != "paragraph":
                    continue
                sid = style.get(_w_attr("styleId"), "")
                if not sid:
                    continue
                style_ids.add(sid)
                name_node = style.find("w:name", W)
                name = name_node.get(_w_attr("val"), sid) if name_node is not None else sid
                style_name_by_id[sid] = name
                style_name_to_id[name.casefold()] = sid
        for name in zf.namelist():
            if not name.endswith(".xml"):
                continue
            if not (
                name == "word/document.xml"
                or name.startswith("word/header")
                or name.startswith("word/footer")
                or name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}
            ):
                continue
            try:
                root = ET.fromstring(zf.read(name))
            except Exception:
                continue
            for pstyle in root.findall(".//w:pStyle", W):
                sid = pstyle.get(_w_attr("val"), "")
                if sid:
                    usage[sid] += 1
    for sid, count in usage.items():
        if sid not in style_ids:
            unknown[sid] = count
    return {
        "style_ids": style_ids,
        "style_name_by_id": style_name_by_id,
        "style_name_to_id": style_name_to_id,
        "usage": usage,
        "unknown": unknown,
    }


def _resolve_style_id(
    style_ids: set[str],
    style_name_to_id: dict[str, str],
    preferred_ids: list[str],
    preferred_names: list[str],
    fallback: str,
) -> str:
    for sid in preferred_ids:
        if sid in style_ids:
            return sid
    for name in preferred_names:
        sid = style_name_to_id.get(name.casefold())
        if sid and sid in style_ids:
            return sid
    if fallback in style_ids:
        return fallback
    return next(iter(style_ids), fallback)


def build_style_remap(style_info: dict[str, object]) -> dict[str, str]:
    style_ids = style_info.get("style_ids", set())
    style_name_to_id = style_info.get("style_name_to_id", {})
    if not isinstance(style_ids, set) or not style_ids:
        return {}
    if not isinstance(style_name_to_id, dict):
        style_name_to_id = {}
    normal = _resolve_style_id(style_ids, style_name_to_id, ["Normal", "a"], ["normal"], "Normal")
    caption = _resolve_style_id(style_ids, style_name_to_id, ["Caption", "ImageCaption", "ae"], ["caption"], normal)
    h1 = _resolve_style_id(style_ids, style_name_to_id, ["Heading1", "1"], ["heading 1"], normal)
    h2 = _resolve_style_id(style_ids, style_name_to_id, ["Heading2", "2"], ["heading 2"], normal)
    h3 = _resolve_style_id(style_ids, style_name_to_id, ["Heading3", "3"], ["heading 3"], normal)
    h4 = _resolve_style_id(style_ids, style_name_to_id, ["Heading4", "4"], ["heading 4"], h3)
    return {
        "FirstParagraph": normal,
        "First Paragraph": normal,
        "BodyText": normal,
        "Body Text": normal,
        "Compact": normal,
        "BlockText": normal,
        "Block Text": normal,
        "Bibliography": normal,
        "ImageCaption": caption,
        "Image Caption": caption,
        "Caption": caption,
        "Heading1": h1,
        "Heading2": h2,
        "Heading3": h3,
        "Heading4": h4,
    }


def _normalize_with_python_docx(docx_path: Path, remap: dict[str, str]) -> int:
    if DocxDocument is None:
        return 0
    doc = DocxDocument(str(docx_path))
    changed = 0
    for para in doc.paragraphs:
        ppr = para._p.pPr
        if ppr is None or ppr.pStyle is None:
            continue
        sid = ppr.pStyle.val
        if sid in remap:
            ppr.pStyle.val = remap[sid]
            changed += 1
    if changed:
        doc.save(str(docx_path))
    return changed


def _rewrite_docx_styles_xml(docx_path: Path, remap: dict[str, str]) -> Counter[str]:
    replaced: Counter[str] = Counter()
    with NamedTemporaryFile(delete=False, suffix=".docx", dir=str(docx_path.parent)) as tmp:
        tmp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                    try:
                        root = ET.fromstring(data)
                        changed = False
                        for pstyle in root.findall(".//w:pStyle", W):
                            sid = pstyle.get(_w_attr("val"), "")
                            if sid in remap:
                                pstyle.set(_w_attr("val"), remap[sid])
                                replaced[sid] += 1
                                changed = True
                        if changed:
                            data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                    except Exception:
                        pass
                zout.writestr(item, data)
        tmp_path.replace(docx_path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)
    return replaced


def normalize_docx_styles(docx_path: Path) -> tuple[dict[str, object], dict[str, object], dict[str, str]]:
    before = analyze_docx_styles(docx_path)
    remap = build_style_remap(before)
    if remap:
        _normalize_with_python_docx(docx_path, remap)
        _rewrite_docx_styles_xml(docx_path, remap)
    after = analyze_docx_styles(docx_path)
    return before, after, remap


def _counter_lines(counter: Counter[str]) -> list[str]:
    if not counter:
        return ["- none"]
    return [f"- `{key}`: `{value}`" for key, value in sorted(counter.items())]


def write_quality_report(
    *,
    report_path: Path,
    docx_path: Path,
    markdown_path: Path,
    reference_doc: Path | None,
    state: DocxExportState,
    pandoc_version: str,
    style_before: dict[str, object],
    style_after: dict[str, object],
    style_remap: dict[str, str],
    fallback_used: bool,
) -> None:
    before_usage = style_before.get("usage", Counter())
    after_usage = style_after.get("usage", Counter())
    before_unknown = style_before.get("unknown", Counter())
    after_unknown = style_after.get("unknown", Counter())
    lines = [
        "# Thesis DOCX Quality Report",
        "",
        f"- DOCX: `{docx_path}`",
        f"- Markdown: `{markdown_path}`",
        f"- Reference doc: `{reference_doc if reference_doc is not None else 'Pandoc default'}`",
        f"- Pandoc: `{pandoc_version}`",
        f"- Conversion fallback used: `{'yes' if fallback_used else 'no'}`",
        f"- included source files: `{len(state.included_sources)}`",
        "",
        "## Included Sources",
        "",
    ]
    for source in state.included_sources:
        lines.append(f"- `{source}`")
    lines.extend(["", "## Missing Sources", ""])
    lines.extend([f"- `{item}`" for item in state.missing_sources] or ["- none"])
    lines.extend(["", "## Skipped Setup Files", ""])
    lines.extend([f"- `{item}`" for item in state.skipped_files] or ["- none"])
    lines.extend(["", "## Missing Assets", ""])
    lines.extend([f"- `{item}`" for item in state.missing_assets] or ["- none"])
    lines.extend(["", "## Unsupported Objects", ""])
    lines.extend(_counter_lines(state.unsupported_counts))
    lines.extend(["", "## Heading Counts", ""])
    if state.heading_counts:
        for level, count in sorted(state.heading_counts.items()):
            lines.append(f"- H{level}: `{count}`")
    else:
        lines.append("- none")
    lines.extend(["", "## Bibliography", ""])
    lines.extend([f"- `{item}`" for item in state.bibliography_files] or ["- none detected"])
    lines.extend(["", "## Style Usage Before Normalization", ""])
    lines.extend(_counter_lines(before_usage if isinstance(before_usage, Counter) else Counter(before_usage)))
    lines.extend(["", "## Style Usage After Normalization", ""])
    lines.extend(_counter_lines(after_usage if isinstance(after_usage, Counter) else Counter(after_usage)))
    lines.extend(["", "## Unknown Styles", ""])
    lines.append(f"- before: `{sum(before_unknown.values()) if isinstance(before_unknown, Counter) else 0}`")
    lines.append(f"- after: `{sum(after_unknown.values()) if isinstance(after_unknown, Counter) else 0}`")
    lines.extend(["", "## Style Remap", ""])
    lines.extend([f"- `{src}` -> `{dst}`" for src, dst in sorted(style_remap.items())] or ["- none"])
    lines.extend(["", "## Warnings", ""])
    lines.extend([f"- {item}" for item in state.warnings] or ["- none"])
    lines.extend(
        [
            "",
            "## Manual Actions",
            "",
            "- update TOC fields in Word.",
            "- inspect every placeholder for complex tables, algorithms, code listings, TikZ, or PDF-only objects.",
            "- refresh references and bibliography formatting if the school template requires exact Word styles.",
            "- verify cover pages and front matter manually against the official school Word template.",
        ]
    )
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _resolve_output(
    project_dir: Path,
    output: Path | None,
    tex_stem: str,
    *,
    allow_external_output: bool = False,
) -> Path:
    if output is None:
        return project_dir / f"{tex_stem}.docx"
    raw = output.expanduser()
    if raw.is_absolute():
        resolved = raw.resolve()
    else:
        resolved = (project_dir / raw).resolve()
    if not allow_external_output and not _is_within(resolved, project_dir):
        raise DocxExportError(
            "DOCX 输出路径默认必须位于项目目录内；如确需写到项目外，请显式使用 --allow-external-output。"
        )
    return resolved


def export_docx_project(
    project_dir: Path,
    tex_file: str = "main.tex",
    output: Path | None = None,
    reference_doc: Path | None = None,
    keep_markdown: bool = False,
    skip_style_normalization: bool = False,
    allow_external_output: bool = False,
) -> Path:
    project_dir = project_dir.expanduser().resolve()
    tex_path = (project_dir / tex_file).resolve()
    if not tex_path.exists():
        raise FileNotFoundError(f"TeX 主文件不存在：{tex_path}")
    if tex_path.parent != project_dir:
        raise DocxExportError("DOCX 导出当前仅支持项目根目录下的主 TeX 文件。")

    tex_stem = tex_path.stem
    cache_dir = project_dir / CACHE_DIRNAME / DOCX_CACHE_DIRNAME
    if cache_dir.exists():
        shutil.rmtree(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    state = DocxExportState(cache_dir)

    collect_tex_sources(tex_path, project_dir, state)
    markdown = render_markdown(project_dir, tex_path, state)
    markdown_path = cache_dir / f"{tex_stem}.md"
    markdown_path.write_text(markdown, encoding="utf-8")
    if keep_markdown:
        shutil.copy2(markdown_path, project_dir / f"{tex_stem}.docx.md")

    resolved_reference_doc = discover_reference_doc(project_dir, reference_doc)
    if resolved_reference_doc is None:
        state.add_warning("No reference doc found; Pandoc default DOCX styles were used.")
    docx_path = _resolve_output(
        project_dir,
        output,
        tex_stem,
        allow_external_output=allow_external_output,
    )
    pandoc_version = get_pandoc_version()
    build_docx_from_markdown(
        markdown_path,
        docx_path,
        project_dir=project_dir,
        reference_doc=resolved_reference_doc,
        bibliography=state.bibliography_files,
        state=state,
    )

    if skip_style_normalization:
        style_before = analyze_docx_styles(docx_path)
        style_after = style_before
        style_remap: dict[str, str] = {}
    else:
        style_before, style_after, style_remap = normalize_docx_styles(docx_path)

    report_path = cache_dir / f"{tex_stem}_docx_quality_report.md"
    write_quality_report(
        report_path=report_path,
        docx_path=docx_path,
        markdown_path=markdown_path,
        reference_doc=resolved_reference_doc,
        state=state,
        pandoc_version=pandoc_version,
        style_before=style_before,
        style_after=style_after,
        style_remap=style_remap,
        fallback_used=state.fallback_used,
    )
    print(f"✓ DOCX generated: {docx_path}")
    print(f"✓ Markdown source: {markdown_path}")
    print(f"✓ Quality report: {report_path}")
    return docx_path
